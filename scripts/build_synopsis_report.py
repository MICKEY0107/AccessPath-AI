from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "archive" / "AccessPath_AI_Synopsis_Report.docx"

TABLE_FILL = "FFFFFF"
TABLE_BORDER = "000000"
TEXT = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), TABLE_BORDER)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        cell.width = Inches(widths[idx])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, TABLE_FILL)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.color.rgb = TEXT
        run.font.size = Pt(9.5)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for idx, text in enumerate(values):
            cell = row.cells[idx]
            cell.width = Inches(widths[idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_shading(cell, TABLE_FILL)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.font.size = Pt(9.2)
            run.font.color.rgb = TEXT
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True


def add_paragraph(doc: Document, text: str, *, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.08
        p.add_run("•\t")
        p.add_run(item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.08
        p.add_run(f"{index}.\t")
        p.add_run(item)


def add_references(doc: Document, references: list[str]) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, citation in enumerate(references, start=1):
        row = table.add_row()
        prevent_row_split(row)
        number_cell, citation_cell = row.cells
        number_cell.width = Inches(0.35)
        citation_cell.width = Inches(6.45)
        set_cell_margins(number_cell, top=30, start=0, bottom=50, end=40)
        set_cell_margins(citation_cell, top=30, start=0, bottom=50, end=0)
        number_p = number_cell.paragraphs[0]
        number_p.paragraph_format.space_after = Pt(0)
        number_run = number_p.add_run(f"[{index}]")
        number_run.font.size = Pt(9.2)
        citation_p = citation_cell.paragraphs[0]
        citation_p.paragraph_format.space_after = Pt(0)
        citation_p.paragraph_format.line_spacing = 1.02
        citation_run = citation_p.add_run(citation)
        citation_run.font.size = Pt(9.2)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    normal.font.color.rgb = TEXT
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = TEXT
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = TEXT
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)


def build() -> None:
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    add_page_number(section.footer.paragraphs[0])

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(92)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.style = doc.styles["Title"]
    p.add_run("AccessPath AI")
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(12)
    r = p2.add_run("Evidence Driven Keyboard Accessibility Journey Evaluation")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = TEXT
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(28)
    r = p3.add_run("Project Synopsis for Midterm Review")
    r.font.size = Pt(14)
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(8)
    p4.add_run("29 September 2026 | 12:00 PM to 2:00 PM")
    p5 = doc.add_paragraph()
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p5.paragraph_format.space_before = Pt(80)
    p5.add_run("Prepared by the AccessPath AI project team").italic = True
    doc.add_page_break()

    add_heading(doc, "Synopsis Overview", 1)
    add_paragraph(
        doc,
        "AccessPath AI is a research prototype that asks a practical accessibility question: can a person following a declared interaction profile complete an important multi-step web task? The midterm prototype focuses on motor accessibility through keyboard-only input. It evaluates a controlled appointment-booking journey, stores step-by-step evidence, and distinguishes successful completion from a genuine keyboard barrier. An AI model may suggest the next action, but deterministic policy, browser observation, and completion rules remain authoritative.",
    )
    add_heading(doc, "Abstract", 2)
    add_paragraph(
        doc,
        "Automated accessibility tools are valuable for detecting many page-level problems, but they do not by themselves prove that a complete user journey is operable. This project proposes a bounded, evidence-driven evaluator for keyboard-only journeys. A planner receives a limited observation of the current browser state and proposes one permitted keyboard action. A deterministic policy validates the proposal before Playwright performs it. The system records focus movement, actions, screenshots, policy decisions, and the final outcome. Evaluation uses two controlled versions of an appointment-booking journey: an accessible reference and a deliberately broken keyboard-barrier control. The expected contribution is a safe and reproducible bridge between static accessibility checking and task-level evaluation. The current prototype includes a Django dashboard, background worker, persistent evidence store, deterministic planner, and configurable live-model planners. Initial live evidence shows one completed accessible booking and one correctly blocked broken journey without deterministic fallback. The work is not a claim of complete WCAG conformance; it is a focused study of keyboard journey completion, evidence quality, safety, and repeatability.",
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(9)
    p.add_run("Keywords: ").bold = True
    p.add_run("web accessibility, keyboard accessibility, WCAG, large language models, browser automation, task completion, evidence-driven testing")

    add_heading(doc, "Contents", 2)
    contents = [
        "1 Introduction and Problem Definition",
        "2 Literature Review",
        "3 Research Gap",
        "4 Aim, Objectives and Research Questions",
        "5 Proposed Methodology",
        "6 Current Development Status",
        "7 Expected Outcomes and Evaluation",
        "8 Limitations and Risk Management",
        "9 Work Plan",
        "10 Conclusion",
        "References",
    ]
    for item in contents:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)
    doc.add_page_break()

    add_heading(doc, "1 Introduction and Problem Definition", 1)
    add_paragraph(
        doc,
        "Web accessibility enables people with different abilities to perceive, understand, navigate, and interact with digital services. WCAG 2.2 provides a common standard covering a wide range of accessibility needs, while WCAG-EM provides a structured process for defining scope, selecting representative content, evaluating it, and reporting findings [1][2]. For users who cannot rely on a mouse, keyboard operation is essential. A control that cannot receive focus or be activated can prevent completion of an appointment, form, payment, or account task even when the page appears visually correct.",
    )
    add_paragraph(
        doc,
        "Many existing tools scan the current page or Document Object Model and report rule violations. This is useful, but the result is usually page-centred. A multi-step journey is different: each action changes the interface, focus may be lost, validation messages may appear, and a later control may become unreachable. The main project problem is therefore the gap between finding isolated accessibility issues and verifying whether a user following a specific interaction profile can complete a declared task.",
    )
    add_paragraph(
        doc,
        "Problem definition: Existing rule-based and AI-assisted accessibility approaches do not consistently provide safe, reproducible, journey-level evidence showing whether a keyboard-only user can complete a multi-step task. A solution must combine flexible planning with deterministic execution control, observable completion rules, and evidence that a reviewer can inspect.",
        bold_lead="Problem definition:",
    )
    add_heading(doc, "1.1 Project Scope", 2)
    add_bullets(
        doc,
        [
            "Accessibility profile: motor accessibility through keyboard-only interaction.",
            "Primary journey: a multi-step appointment-booking task.",
            "Evaluation targets: controlled accessible and deliberately broken fixture sites.",
            "Allowed interaction: bounded keyboard actions such as Tab, Shift+Tab, Enter, Space, arrow keys, and approved synthetic text entry.",
            "Output: completed, blocked, inconclusive, safety-stopped, provider error, or environment error, supported by step-level evidence.",
        ],
    )

    add_heading(doc, "2 Literature Review", 1)
    add_heading(doc, "2.1 Web Accessibility Standards", 2)
    add_paragraph(
        doc,
        "WCAG 2.2 organizes web accessibility around perceivable, operable, understandable, and robust content [1]. The keyboard-focused criteria most relevant to this prototype include Keyboard, No Keyboard Trap, Focus Order, Focus Visible, and Focus Not Obscured. These criteria describe necessary properties, but a complete evaluation still requires a defined scope, representative views and states, suitable tools, expert judgment, and clear reporting. WCAG-EM formalizes this broader evaluation process and recommends involving users with disabilities [2]. AccessPath AI supports a narrow keyboard profile; it does not claim complete WCAG conformance.",
    )
    add_heading(doc, "2.2 Traditional Automated Evaluation", 2)
    add_paragraph(
        doc,
        "Rule-based tools such as axe-core inspect rendered interfaces against accessibility rules. They are fast and useful in development pipelines. However, axe-core's own documentation reports average automatic coverage of 57% of WCAG issues and returns uncertain cases for manual review [3]. Playwright's accessibility guidance similarly recommends combining automated checks with manual assessment and inclusive user testing [4]. This establishes that automation supports, but does not replace, broader evaluation.",
    )
    add_paragraph(
        doc,
        "Vigo, Brown, and Conway compared six accessibility evaluation tools and showed the risk of relying only on automated checks [5]. Ara, Sik-Lanyi, and Kelemen reviewed 92 primary studies and found a wide range of methods, tools, and engineering assets, together with continuing research gaps across the evaluation process [6]. Pool compared nine tools on 121 pages and found that each detected issues missed by the others [7]. These studies show that tool outputs are complementary and that no single scanner provides complete evidence.",
    )
    add_heading(doc, "2.3 Keyboard and Journey Level Evaluation", 2)
    add_paragraph(
        doc,
        "Keyboard accessibility is often assessed manually by repeatedly pressing Tab and related keys while observing reachability, activation, focus order, focus visibility, traps, and recovery. This can reveal behavioural barriers that a single-state scan may miss, but manual execution is time-consuming and difficult to repeat consistently. Scripted end-to-end tests improve repeatability, yet they require a human to define the exact path in advance. AccessPath AI investigates whether a bounded planner can choose the path while deterministic rules preserve safety and evidence quality.",
    )
    add_heading(doc, "2.4 AI and LLM Based Evaluation", 2)
    add_paragraph(
        doc,
        "Recent research uses AI for alternative text, compliance assessment, remediation suggestions, accessible code generation, and interface adaptation. Vera-Amaro and Rojano-Caceres reviewed 277 studies and selected 31 key works, identifying growing AI use alongside inconsistent compliance and limited coverage beyond specific code cases [8]. Lopez-Gil and Pereira used LLM-based scripts for three WCAG criteria that normally require manual judgment and reported 87.18% detection across their test cases [9]. This suggests that language models can add semantic judgment where fixed rules struggle, but the evidence is limited to selected criteria and datasets.",
    )
    add_paragraph(
        doc,
        "Aljedaani and Mollik reviewed 38 peer-reviewed studies on LLMs for web accessibility. They found that research was concentrated on text-centric and structurally explicit tasks, evaluation methods varied, and direct participation by people with disabilities was limited [10]. Agentic research is also emerging. Yepes-Parra and Escobar-Velasquez proposed agents that interact through keyboard and screen-reader modalities, explicitly arguing that static analysis cannot capture all real-time interaction problems [11]. This work is close to the direction of AccessPath AI, so the project does not claim that AI agents for accessibility are entirely new.",
    )

    add_heading(doc, "2.5 Literature Comparison", 2)
    add_table(
        doc,
        ["Source", "Main contribution", "Limitation relevant to this project"],
        [
            ["WCAG 2.2 and WCAG-EM", "Standards and structured evaluation process", "Do not themselves execute a user's complete task"],
            ["Vigo et al. 2013", "Measured coverage and correctness of six tools", "Shows risk of relying only on automated scanners"],
            ["Ara et al. 2024", "Review of 92 accessibility-evaluation studies", "Reports diverse tools and continuing evaluation gaps"],
            ["Pool 2023", "Compared nine tools on 121 pages", "Tools were complementary rather than complete"],
            ["Lopez-Gil and Pereira 2025", "LLM checks for three manual WCAG criteria", "Selected criteria; not multi-step task completion"],
            ["Vera-Amaro and Rojano-Caceres 2025", "Review of AI for web accessibility", "Inconsistent compliance and limited non-visual coverage"],
            ["Aljedaani and Mollik 2026", "Review of 38 LLM-accessibility studies", "Concentration on text and explicit structures; varied evaluation"],
            ["Yepes-Parra and Escobar-Velasquez 2025", "Emerging autonomous interaction framework", "Early research with limited standardized benchmarks"],
        ],
        [1.55, 2.5, 2.75],
    )

    add_heading(doc, "3 Research Gap", 1)
    add_paragraph(
        doc,
        "The literature does not support the broad claim that AI has never been used for accessibility. A more defensible gap is the limited evidence on bounded AI agents that must complete a real multi-step journey through a specific accessibility interaction model while deterministic components retain control of safety, execution, outcome classification, and evidence.",
    )
    add_bullets(
        doc,
        [
            "Page-level checks versus task-level outcomes: scanners identify violations, but users experience success or failure across a sequence of states.",
            "Detection versus interaction: many systems analyze HTML or text without discovering and executing a journey path.",
            "Semantic reasoning versus execution safety: LLMs can reason about intent but may return invalid, repeated, or unsafe actions.",
            "Finding reports versus reproducible evidence: developers need the exact state, focused element, action, result, and screenshot that produced a barrier.",
            "Prototype demonstrations versus controlled benchmarks: comparable accessible and broken versions with known ground truth are still uncommon.",
        ],
    )
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    add_paragraph(
        doc,
        "Research gap statement: There is an opportunity for a bounded, evidence-driven framework in which an AI planner selects from permitted keyboard actions, while deterministic browser observation, policy validation, completion oracles, and stored evidence remain authoritative. AccessPath AI studies this combination using controlled accessible and broken appointment-booking journeys.",
        bold_lead="Research gap statement:",
    )

    add_heading(doc, "4 Aim, Objectives and Research Questions", 1)
    add_heading(doc, "4.1 Aim", 2)
    add_paragraph(
        doc,
        "To design and evaluate a safe, repeatable system that determines whether a declared multi-step web journey can be completed using a keyboard-only accessibility profile, with AI-assisted planning and deterministic evidence-based outcomes.",
    )
    add_heading(doc, "4.2 Objectives", 2)
    add_numbered(
        doc,
        [
            "Build a controlled accessible journey and a deliberately broken version with known ground truth.",
            "Represent browser observations, keyboard actions, policy decisions, and evidence using strict versioned data contracts.",
            "Allow a deterministic planner and selected live-model planners to propose the next bounded keyboard action.",
            "Validate every proposal through deterministic safety and focus rules before browser execution.",
            "Classify journey outcomes using observable completion and barrier conditions rather than model opinion.",
            "Store step-level evidence that enables reproduction and human review.",
            "Measure task completion, correct blocking, false completion, fallback use, and repeated-run agreement.",
        ],
    )
    add_heading(doc, "4.3 Research Questions", 2)
    add_table(
        doc,
        ["ID", "Research question"],
        [
            ["RQ1", "Can a bounded AI planner complete a multi-step web task using only permitted keyboard interactions?"],
            ["RQ2", "Can journey-level execution identify keyboard barriers that are not adequately represented by a single page scan?"],
            ["RQ3", "Does deterministic policy validation prevent invalid or unsafe model actions from reaching the browser?"],
            ["RQ4", "How reliably can the system distinguish completion, accessibility blocking, provider failure, environment failure, and inconclusive results?"],
            ["RQ5", "Does step-level evidence improve reproducibility and human review of reported barriers?"],
        ],
        [0.8, 6.0],
    )

    add_heading(doc, "5 Proposed Methodology", 1)
    add_heading(doc, "5.1 Research Design", 2)
    add_paragraph(
        doc,
        "The study follows a design-and-evaluate approach. First, the team builds a bounded prototype and controlled test environment. Second, the same declared task is executed against accessible and broken journey variants. Third, outcomes are compared with known ground truth and reviewed through stored evidence. The study treats AI as a planner, not as the authority for safety or completion.",
    )
    add_heading(doc, "5.2 Dataset and Test Environment", 2)
    add_paragraph(
        doc,
        "The Phase 1 dataset is a versioned set of controlled appointment-booking fixtures. The accessible fixture supports keyboard completion. The broken control contains one declared barrier: the appointment time-slot control cannot be reached or operated through the permitted keyboard path. All form values are synthetic. No real patient or personal data is required. Controlled fixtures provide known ground truth and make repeated runs comparable.",
    )
    add_heading(doc, "5.3 System Flow", 2)
    add_table(
        doc,
        ["Stage", "Responsibility"],
        [
            ["1 Observe", "Capture the approved URL, visible text, focused element, editable state, progress, and remaining budget."],
            ["2 Propose", "The selected planner proposes one structured keyboard action."],
            ["3 Validate", "Deterministic policy checks action type, origin, target, focus state, typing permission, and budgets."],
            ["4 Execute", "Playwright performs only the approved keyboard action."],
            ["5 Record", "Store before and after states, action result, screenshots, hashes, and planner metadata."],
            ["6 Decide", "A deterministic oracle checks declared completion or barrier conditions."],
            ["7 Repeat or stop", "Continue within limits, or finish with an explicit outcome and reason."],
        ],
        [1.15, 5.65],
    )
    add_heading(doc, "5.4 Tools and Technologies", 2)
    add_table(
        doc,
        ["Technology", "Use"],
        [
            ["Python and Django", "Web application, data models, run workflow, validation, and reporting UI"],
            ["PostgreSQL", "Journeys, run state, steps, provider metadata, and artifact records"],
            ["Redis and Celery", "Queued background execution with a single-concurrency midterm worker"],
            ["Playwright", "Keyboard events, browser observation, navigation, and screenshots"],
            ["Pydantic", "Strict observation, action, policy, and evidence contracts"],
            ["Gemini, OpenRouter, and experimental Groq", "Configurable live action proposals; each provider is recorded separately"],
            ["Deterministic planner", "Repeatable fixture execution and contingency demonstration"],
            ["Docker Compose", "Reproducible local database, queue, fixture, web, and worker services"],
            ["pytest, Ruff, and Pyright", "Automated tests, linting, and static type checks"],
        ],
        [2.1, 4.7],
    )
    add_heading(doc, "5.5 Safety Controls", 2)
    add_bullets(
        doc,
        [
            "Allowlisted fixture origins and bounded navigation scope.",
            "Small keyboard action schema with no arbitrary scripts or direct mouse clicks.",
            "Policy approval before every action.",
            "Synthetic input values only.",
            "Call, step, timeout, retry, and rejection limits.",
            "Visible provider and fallback metadata; deterministic evidence is never presented as live-model evidence.",
            "Process-wide pacing between requests to the same provider in the current worker process.",
        ],
    )

    add_heading(doc, "6 Current Development Status", 1)
    add_paragraph(
        doc,
        "The midterm vertical slice is implemented and in final evidence review. The local Docker environment includes PostgreSQL, Redis, a controlled fixture server, the Django web application, and a Celery worker. The dashboard allows a reviewer to choose a journey and planner, start a run, and inspect the persisted result. Run pages show outcomes, step details, provider metadata, screenshots, and evidence hashes.",
    )
    add_table(
        doc,
        ["Component", "Current status"],
        [
            ["Controlled accessible and broken fixtures", "Implemented and manually exercised"],
            ["Keyboard executor and completion oracle", "Implemented with persisted evidence"],
            ["Deterministic planner", "Working and repeatable"],
            ["Gemini planner", "Implemented; produced stored live completion and blocking evidence; quota remains variable"],
            ["OpenRouter planner", "Implemented; free routing showed upstream overload and incomplete output"],
            ["Groq planner", "Experimental addition; not included in verified midterm evidence until tests and preflight are complete"],
            ["Provider dropdown", "Implemented; provider and model are stored with each run"],
            ["Documentation", "Architecture, handoff, demo runbook, and presentation outline prepared"],
        ],
        [2.65, 4.15],
    )
    add_heading(doc, "6.1 Current Evidence", 2)
    add_paragraph(
        doc,
        "Two persisted Gemini runs provide the clearest final individual midterm proofs. Run af03b14a-6700-4b91-a437-b8e3ada1de83 completed the accessible booking in 29 planner calls without fallback. Run cf0e9180-6cd1-4d08-a7f1-972eba776d56 classified the broken fixture as blocked in 13 calls without fallback, with terminal reason KEYBOARD_TARGET_NOT_REACHABLE. Historical live evidence also includes three accessible completions and five broken-control runs with zero false completions. However, the uninterrupted post-fix provider-pure 3+5 matrix remains pending and should not be presented as a completed final gate.",
    )
    add_paragraph(
        doc,
        "The latest documented quality check reports 48 automated tests passing together with Ruff, Pyright, Django system, and migration-drift checks. Headed-browser review also confirmed the provider dropdown and keyboard-accessible dashboard. Groq remains an experimental planner and is not included in the verified midterm evidence.",
    )

    add_heading(doc, "7 Expected Outcomes and Evaluation", 1)
    add_heading(doc, "7.1 Outcome Classes", 2)
    add_bullets(
        doc,
        [
            "Completed: the declared confirmation condition is observed.",
            "Blocked: a defined keyboard barrier prevents progress and is supported by evidence.",
            "Inconclusive: the run ends without enough evidence for completion or a confirmed barrier.",
            "Safety stopped: a proposal or navigation condition violates policy.",
            "Provider or environment error: an external model or local service prevents evaluation.",
        ],
    )
    add_heading(doc, "7.2 Evaluation Measures", 2)
    add_table(
        doc,
        ["Measure", "Interpretation"],
        [
            ["Task completion rate", "Percentage of accessible runs that reach the declared completion state"],
            ["Correct block rate", "Percentage of broken runs correctly classified as blocked"],
            ["False completion rate", "Broken runs incorrectly reported as completed; target is zero"],
            ["False block rate", "Accessible runs incorrectly reported as blocked"],
            ["Fallback rate", "Runs in which live planning switches to deterministic contingency"],
            ["Repeated-run agreement", "Consistency of outcomes across repeated runs"],
            ["Evidence completeness", "Availability of action, focus, result, screenshot, reason, and provider metadata"],
            ["Efficiency", "Calls, executed steps, elapsed time, and provider failures"],
        ],
        [2.15, 4.65],
    )
    add_paragraph(
        doc,
        "The Phase 1 exit target is a provider-pure matrix of three accessible completions and five broken runs with zero false completions and no deterministic fallback counted as live-model evidence. Provider results must not be mixed. The team must first confirm whether the academic requirement specifically demands Gemini or accepts any fixed, preflight-verified live provider.",
    )

    add_heading(doc, "8 Limitations and Risk Management", 1)
    add_table(
        doc,
        ["Limitation or risk", "Mitigation"],
        [
            ["Provider quota and temporary outages", "Persist successful evidence, use bounded retries and pacing, keep deterministic contingency clearly labelled"],
            ["Changing free-router models", "Use a fixed preflight-verified model for final evidence"],
            ["LLM invalid or unsafe actions", "Strict action schema, server validation, deterministic policy, and bounded budgets"],
            ["Overstating accessibility coverage", "Report a keyboard-profile result, not full WCAG or legal conformance"],
            ["Limited benchmark", "Use controlled ground truth now; expand journey families and barriers later"],
            ["No user study yet", "Treat the prototype as engineering evidence and add reviewers and participants in later evaluation"],
            ["Local single-user deployment", "Defer authentication, authorization, cloud storage, CI/CD, and production hardening"],
        ],
        [2.35, 4.45],
    )

    add_heading(doc, "9 Work Plan", 1)
    add_table(
        doc,
        ["Stage", "Planned work", "Expected output"],
        [
            ["Before midterm", "Stabilize or exclude experimental Groq path; confirm provider evidence rule; final verification and rehearsal", "Clean demo revision and honest limitations"],
            ["Phase 2", "Authentication, projects, authorization, durable jobs, cloud-compatible artifacts, security controls", "Reliable multi-user platform foundation"],
            ["Phase 3", "More keyboard barriers, journey families, replay, comparison, and benchmark measurements", "Broader keyboard evaluation coverage"],
            ["Phase 4", "Safe sandbox transactions, report export, secret references, cleanup verification", "Controlled complete-journey workflows"],
            ["Phase 5", "Reviewed benchmark, precision and recall analysis, WCAG coverage register, user involvement", "Academic evaluation and compliance traceability"],
            ["Phase 6", "Staging deployment, CI/CD, observability, threat model, backup and recovery", "Final hardened demonstration"],
        ],
        [1.15, 3.55, 2.1],
    )

    add_heading(doc, "10 Conclusion", 1)
    doc.paragraphs[-1].paragraph_format.page_break_before = True
    add_paragraph(
        doc,
        "AccessPath AI addresses a clear gap between page-level accessibility findings and evidence that a user can complete a multi-step task. The proposed approach combines the flexibility of AI planning with deterministic control over browser actions, completion decisions, and evidence. The midterm prototype demonstrates the approach for keyboard-only appointment booking using accessible and deliberately broken fixtures. Its contribution is not a claim that AI can certify accessibility. It is a practical, auditable method for studying journey completion, barriers, safety, and reproducibility. The next stages will broaden the benchmark, improve platform reliability, include human review, and measure accuracy across more journeys and accessibility profiles.",
    )

    add_heading(doc, "References", 1)
    doc.paragraphs[-1].paragraph_format.page_break_before = True
    references = [
        "World Wide Web Consortium. (2024). Web Content Accessibility Guidelines (WCAG) 2.2. https://www.w3.org/TR/WCAG22/",
        "World Wide Web Consortium. (2026). WCAG Evaluation Methodology (WCAG-EM) 2.0. https://www.w3.org/TR/WCAG-EM/",
        "Deque Systems. (2026). axe-core: Accessibility engine for automated Web UI testing. https://github.com/dequelabs/axe-core",
        "Microsoft. (2026). Accessibility testing with Playwright. https://playwright.dev/docs/accessibility-testing",
        "Vigo, M., Brown, J., & Conway, V. (2013). Benchmarking web accessibility evaluation tools: Measuring the harm of sole reliance on automated tests. W4A 2013. https://doi.org/10.1145/2461121.2461124",
        "Ara, J., Sik-Lanyi, C., & Kelemen, A. (2024). Accessibility engineering in web evaluation process: A systematic literature review. Universal Access in the Information Society, 23, 653-686. https://doi.org/10.1007/s10209-023-00967-2",
        "Pool, J. R. (2023). Accessibility Metatesting: Comparing Nine Testing Tools. W4A 2023. https://doi.org/10.1145/3587281.3587282",
        "Vera-Amaro, G., & Rojano-Caceres, J. R. (2025). Towards accessible website design through artificial intelligence: A systematic literature review. Information and Software Technology, 186, 107821. https://doi.org/10.1016/j.infsof.2025.107821",
        "Lopez-Gil, J. M., & Pereira, J. (2025). Turning manual web accessibility success criteria into automatic: An LLM-based approach. Universal Access in the Information Society, 24, 837-852. https://doi.org/10.1007/s10209-024-01108-z",
        "Aljedaani, W., & Mollik, R. H. (2026). Large Language Models for Web Accessibility: A Systematic Literature Review. W4A 2026. https://doi.org/10.1145/3800424.3800452",
        "Yepes-Parra, J. D., & Escobar-Velasquez, C. (2025). Autonomous Agents for Accessibility: Simulating Visual Impairments in Web Interfaces. ASE 2025, 3953-3957. https://doi.org/10.1109/ASE63991.2025.00349",
        "Microsoft. (2026). Keyboard API. Playwright documentation. https://playwright.dev/docs/api/class-keyboard",
    ]
    add_references(doc, references)

    # Prevent stray landscape/default section changes if Word adds one later.
    for sec in doc.sections:
        if sec.start_type == WD_SECTION.NEW_PAGE:
            sec.page_width = Inches(8.5)
            sec.page_height = Inches(11)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
