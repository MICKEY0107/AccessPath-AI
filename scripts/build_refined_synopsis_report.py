from __future__ import annotations

import math
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "AccessPath_AI_Synopsis_Report_Refined.docx"
ARTIFACT_DIR = ROOT / ".artifacts" / "synopsis-refined"
DIAGRAM = ARTIFACT_DIR / "accesspath_architecture.png"

NAVY = "1F4E78"
BLUE = "DCEAF7"
PALE_BLUE = "F4F8FC"
LIGHT_GRAY = "D9D9D9"
MID_GRAY = "6B7280"
TEXT_HEX = "111827"
WHITE = "FFFFFF"
TEXT = RGBColor(17, 24, 39)


def font_path(name: str) -> str:
    path = Path("C:/Windows/Fonts") / name
    if not path.exists():
        raise FileNotFoundError(path)
    return str(path)


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(font_path("segoeuib.ttf" if bold else "segoeui.ttf"), size)


def wrapped_lines(draw: ImageDraw.ImageDraw, text: str, font, max_width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    line = ""
    for word in words:
        candidate = word if not line else f"{line} {word}"
        if draw.textlength(candidate, font=font) <= max_width:
            line = candidate
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)
    return lines


def draw_text_block(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    title: str,
    detail: str,
    *,
    title_color: str = TEXT_HEX,
    detail_color: str = MID_GRAY,
    title_size: int = 38,
    detail_size: int = 28,
) -> None:
    x1, y1, x2, y2 = box
    title_font = load_font(title_size, bold=True)
    detail_font = load_font(detail_size)
    title_lines = wrapped_lines(draw, title, title_font, x2 - x1 - 50)
    detail_lines = wrapped_lines(draw, detail, detail_font, x2 - x1 - 55)
    title_height = len(title_lines) * (title_size + 8)
    detail_height = len(detail_lines) * (detail_size + 8)
    gap = 14 if detail else 0
    y = y1 + max(20, ((y2 - y1) - title_height - detail_height - gap) // 2)
    for line in title_lines:
        width = draw.textlength(line, font=title_font)
        draw.text(((x1 + x2 - width) / 2, y), line, font=title_font, fill=f"#{title_color}")
        y += title_size + 8
    y += gap
    for line in detail_lines:
        width = draw.textlength(line, font=detail_font)
        draw.text(((x1 + x2 - width) / 2, y), line, font=detail_font, fill=f"#{detail_color}")
        y += detail_size + 8


def rounded_box(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    title: str,
    detail: str,
    *,
    fill: str = WHITE,
    outline: str = LIGHT_GRAY,
    width: int = 4,
    title_size: int = 38,
    detail_size: int = 28,
) -> None:
    draw.rounded_rectangle(box, radius=28, fill=f"#{fill}", outline=f"#{outline}", width=width)
    draw_text_block(
        draw,
        box,
        title,
        detail,
        title_size=title_size,
        detail_size=detail_size,
    )


def arrow(
    draw: ImageDraw.ImageDraw,
    start: tuple[int, int],
    end: tuple[int, int],
    *,
    color: str = NAVY,
    width: int = 7,
    head: int = 22,
) -> None:
    draw.line([start, end], fill=f"#{color}", width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    left = (
        end[0] - head * math.cos(angle - math.pi / 6),
        end[1] - head * math.sin(angle - math.pi / 6),
    )
    right = (
        end[0] - head * math.cos(angle + math.pi / 6),
        end[1] - head * math.sin(angle + math.pi / 6),
    )
    draw.polygon([end, left, right], fill=f"#{color}")


def poly_arrow(
    draw: ImageDraw.ImageDraw,
    points: list[tuple[int, int]],
    *,
    color: str = NAVY,
    width: int = 7,
) -> None:
    draw.line(points, fill=f"#{color}", width=width, joint="curve")
    arrow(draw, points[-2], points[-1], color=color, width=width)


def build_architecture_diagram(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (2600, 1500), "#FFFFFF")
    draw = ImageDraw.Draw(image)

    title_font = load_font(60, bold=True)
    subtitle_font = load_font(30)
    draw.text((90, 55), "AccessPath AI system architecture", font=title_font, fill=f"#{TEXT_HEX}")
    draw.text(
        (92, 135),
        "AI suggests an action; fixed rules control execution, results, and evidence.",
        font=subtitle_font,
        fill=f"#{MID_GRAY}",
    )

    reviewer = (90, 245, 430, 445)
    dashboard = (535, 245, 995, 445)
    worker = (1100, 245, 1640, 445)
    evidence = (1760, 245, 2500, 445)
    rounded_box(draw, reviewer, "Reviewer", "Starts a run and checks the result", fill=PALE_BLUE, outline=NAVY)
    rounded_box(draw, dashboard, "Django dashboard", "Selects the journey and planner", fill=PALE_BLUE, outline=NAVY)
    rounded_box(draw, worker, "Redis queue and Celery worker", "Runs the evaluation in the background", fill=PALE_BLUE, outline=NAVY)
    rounded_box(draw, evidence, "Run record and evidence", "PostgreSQL data and screenshot files", fill=PALE_BLUE, outline=NAVY)
    arrow(draw, (430, 345), (535, 345))
    arrow(draw, (995, 345), (1100, 345))
    arrow(draw, (1640, 345), (1760, 345))

    loop = (90, 570, 2500, 1085)
    draw.rounded_rectangle(loop, radius=34, fill="#F8FAFC", outline=f"#{LIGHT_GRAY}", width=4)
    loop_title = load_font(36, bold=True)
    loop_note = load_font(27)
    draw.text((140, 600), "Controlled evaluation loop", font=loop_title, fill=f"#{TEXT_HEX}")
    draw.text(
        (140, 650),
        "The loop repeats until the task completes, a keyboard barrier is confirmed, or a safety limit is reached.",
        font=loop_note,
        fill=f"#{MID_GRAY}",
    )

    boxes = [
        (145, 750, 545, 960, "Observe page", "URL, text, focus, and progress"),
        (625, 750, 1025, 960, "Planner suggests one action", "AI model or fixed plan"),
        (1105, 750, 1505, 960, "Safety check", "Only approved keyboard input continues"),
        (1585, 750, 1985, 960, "Playwright executes", "Sends keyboard input"),
        (2065, 750, 2465, 960, "Outcome rules", "Completion, barrier, or labelled issue"),
    ]
    for x1, y1, x2, y2, title, detail in boxes:
        rounded_box(
            draw,
            (x1, y1, x2, y2),
            title,
            detail,
            fill=WHITE,
            outline=NAVY,
            title_size=42,
            detail_size=32,
        )
    for left, right in zip(boxes, boxes[1:], strict=False):
        arrow(draw, (left[2], 855), (right[0], 855))

    poly_arrow(draw, [(2265, 960), (2265, 1030), (345, 1030), (345, 965)], color=MID_GRAY, width=5)

    poly_arrow(
        draw,
        [(1370, 445), (1370, 500), (50, 500), (50, 700), (345, 700), (345, 750)],
        width=6,
    )

    test_site = (1260, 1180, 1880, 1390)
    store = (1950, 1180, 2500, 1390)
    rounded_box(
        draw,
        test_site,
        "Controlled test site",
        "Accessible version or deliberately broken version",
        fill=BLUE,
        outline=NAVY,
        title_size=40,
        detail_size=30,
    )
    rounded_box(
        draw,
        store,
        "Evidence saved after each step",
        "Action, focus, result, screenshot, and provider details",
        fill=BLUE,
        outline=NAVY,
        title_size=36,
        detail_size=29,
    )
    arrow(draw, (1785, 960), (1570, 1180))
    arrow(draw, (1570, 1180), (1785, 965))
    arrow(draw, (2265, 960), (2225, 1180))
    poly_arrow(draw, [(2225, 1180), (2225, 1120), (2480, 1120), (2480, 450)], width=6)

    image.save(path, format="PNG", dpi=(300, 300), optimize=True)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=115, start=135, bottom=115, end=135) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), LIGHT_GRAY)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_cell_width(cell, inches: float) -> None:
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_font(run, name: str = "Aptos") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_font(run)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, end])


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(29)
    title.font.bold = True
    title.font.color.rgb = TEXT
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title_ppr = title._element.get_or_add_pPr()
    border = title_ppr.find(qn("w:pBdr"))
    if border is not None:
        title_ppr.remove(border)

    for style_name, size, before, after in (
        ("Heading 1", 16, 14, 6),
        ("Heading 2", 12.5, 10, 4),
        ("Heading 3", 11, 8, 3),
    ):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = TEXT
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Aptos"
        style.font.size = Pt(10.5)
        style.font.color.rgb = TEXT
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.paragraph_format.space_after = Pt(3)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_paragraph(doc: Document, text: str, *, bold_lead: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
        paragraph.add_run(text[len(bold_lead) :])
    else:
        paragraph.add_run(text)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(item, style="List Bullet")
        paragraph.paragraph_format.left_indent = Inches(0.28)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(item, style="List Number")
        paragraph.paragraph_format.left_indent = Inches(0.32)
        paragraph.paragraph_format.first_line_indent = Inches(-0.22)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    *,
    font_size: float = 9.0,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for index, heading in enumerate(headers):
        cell = header.cells[index]
        set_cell_width(cell, widths[index])
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(heading)
        set_font(run)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(font_size)

    for row_number, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_number % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.03
            if len(value) < 16 and index == 0:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(value)
            set_font(run)
            run.font.size = Pt(font_size)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.space_before = Pt(0)


def add_figure(doc: Document, image_path: Path) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    inline = run.add_picture(str(image_path), width=Inches(6.75))
    doc_pr = inline._inline.docPr
    doc_pr.set("name", "AccessPath AI system architecture")
    doc_pr.set(
        "descr",
        "System flow from reviewer and Django dashboard through the Redis and Celery worker, the controlled keyboard evaluation loop, the test site, and stored evidence.",
    )
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = False
    run = caption.add_run("Figure 1  AccessPath AI system architecture and evidence flow")
    set_font(run)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(75, 85, 99)


def add_references(doc: Document, references: list[str]) -> None:
    for index, citation in enumerate(references, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.42)
        paragraph.paragraph_format.first_line_indent = Inches(-0.42)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.keep_together = True
        number_run = paragraph.add_run(f"[{index}]\t")
        set_font(number_run)
        number_run.font.size = Pt(8.7)
        citation_run = paragraph.add_run(citation)
        set_font(citation_run)
        citation_run.font.size = Pt(8.7)


def build_report() -> None:
    build_architecture_diagram(DIAGRAM)
    doc = Document()
    configure_styles(doc)

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    add_page_number(section.footer.paragraphs[0])

    # Cover page
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(90)
    title.add_run("AccessPath AI Project Synopsis")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(12)
    subtitle.paragraph_format.space_after = Pt(4)
    run = subtitle.add_run("Keyboard Journey Accessibility Evaluation")
    set_font(run, "Aptos Display")
    run.font.size = Pt(17)
    run.bold = True

    review = doc.add_paragraph()
    review.alignment = WD_ALIGN_PARAGRAPH.CENTER
    review.paragraph_format.space_before = Pt(8)
    run = review.add_run("Synopsis for Midterm Review")
    set_font(run)
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(75, 85, 99)

    authors = doc.add_paragraph()
    authors.alignment = WD_ALIGN_PARAGRAPH.CENTER
    authors.paragraph_format.space_before = Pt(82)
    authors.paragraph_format.line_spacing = 1.5
    authors.add_run("Aniket Mishra - 230648\nKrish Bajaj - 230665\nDaksh Choudhary - 230876")

    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_p.paragraph_format.space_before = Pt(42)
    run = date_p.add_run("October 2026")
    set_font(run)
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(75, 85, 99)
    doc.add_page_break()

    add_heading(doc, "Synopsis Overview", 1)
    add_paragraph(
        doc,
        "AccessPath AI tests whether a person who relies on a keyboard can complete a defined web task. The current prototype runs an appointment-booking journey on two controlled sites: one accessible and one intentionally blocked. An AI model may suggest the next key press, but fixed rules decide what is safe to execute and whether the task has completed or failed. Every step is stored for review.",
    )

    add_heading(doc, "Abstract", 2)
    add_paragraph(
        doc,
        "Page scanners can find many accessibility problems, but they do not show whether a user can finish a task that spans several pages and changing interface states. AccessPath AI evaluates that full journey for a keyboard-only user. At each step, the system records a limited view of the page, asks a selected planner for one keyboard action, checks the action against fixed safety rules, and executes it with Playwright only if approved. It then stores focus changes, action results, screenshots, and the final outcome. The study uses two controlled appointment-booking sites with known expected results: an accessible version and a version with an unreachable time-slot control. The prototype includes a Django dashboard, a background worker, persistent run records, a fixed test planner, and optional live AI planners. The project evaluates task completion, barrier detection, safety, and repeatability. It does not claim full WCAG conformance.",
    )

    keywords = doc.add_paragraph()
    keywords.paragraph_format.space_before = Pt(2)
    run = keywords.add_run("Keywords  ")
    set_font(run)
    run.bold = True
    keywords.add_run(
        "web accessibility, keyboard access, WCAG, AI planning, browser automation, task completion, test evidence"
    )

    add_heading(doc, "Contents", 2)
    contents = [
        "1 Introduction and Problem Definition",
        "2 Literature Review",
        "3 Research Gap",
        "4 Aim Objectives and Research Questions",
        "5 Proposed Methodology",
        "6 Expected Outcomes and Evaluation",
        "7 Conclusion",
        "References",
    ]
    for item in contents:
        p = doc.add_paragraph(item)
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.space_after = Pt(1.5)
    doc.add_page_break()

    add_heading(doc, "1 Introduction and Problem Definition", 1)
    add_paragraph(
        doc,
        "Web accessibility allows people with different abilities to use digital services. WCAG 2.2 defines widely used accessibility requirements, and WCAG-EM explains how to scope, evaluate, and report an accessibility review [1][2]. For a person who cannot use a mouse, keyboard access is essential. If focus cannot reach or activate one control, an appointment, form, payment, or account task may become impossible.",
    )
    add_paragraph(
        doc,
        "Most automated tools inspect the current page and report rule failures. That is useful, but a multi-step task changes after every action. Focus can move incorrectly, validation can appear, and a later control can become unreachable. The project therefore studies task completion across the whole journey, not only problems on one page.",
    )
    add_paragraph(
        doc,
        "Problem definition: Current rule-based and AI-assisted tools do not consistently produce safe, repeatable evidence that a keyboard-only user can complete a multi-step task. A useful system must limit what the AI can request, check each action before execution, use observable rules for the final result, and keep evidence that another person can review.",
        bold_lead="Problem definition:",
    )

    add_heading(doc, "1.1 Project Scope", 2)
    add_bullets(
        doc,
        [
            "User profile: motor accessibility through keyboard-only interaction.",
            "Main task: a multi-step appointment-booking journey.",
            "Test targets: controlled accessible and deliberately broken sites.",
            "Permitted input: Tab, Shift+Tab, Enter, Space, arrow keys, and approved synthetic text entry.",
            "Possible results: completed, blocked, inconclusive, safety stopped, provider error, or environment error.",
        ],
    )

    add_heading(doc, "2 Literature Review", 1)
    add_heading(doc, "2.1 Standards and Automated Checks", 2)
    add_paragraph(
        doc,
        "WCAG 2.2 groups accessibility under four principles: perceivable, operable, understandable, and robust [1]. The criteria most relevant here cover keyboard operation, keyboard traps, focus order, focus visibility, and focus that is not hidden. WCAG-EM also requires clear scope, representative states, suitable tools, expert review, and reporting [2]. AccessPath AI addresses only a narrow keyboard profile and cannot certify full conformance.",
    )
    add_paragraph(
        doc,
        "Rule-based tools such as axe-core are fast and useful during development, but Deque reports that axe-core automatically detects an average of 57 percent of WCAG issues [3]. Playwright also recommends combining automated checks with manual assessment and user testing [4]. Comparisons of accessibility tools show that each tool finds issues that others miss [5][6][7].",
    )

    add_heading(doc, "2.2 Keyboard Journeys and AI", 2)
    add_paragraph(
        doc,
        "Manual keyboard testing can reveal unreachable controls, poor focus order, traps, and failed recovery. It is valuable but slow to repeat. Fixed end-to-end scripts are repeatable, yet a person must define every step in advance. This project tests whether a planner can choose the next step while fixed rules continue to control safety and results.",
    )
    add_paragraph(
        doc,
        "Recent studies use artificial intelligence for accessibility checks, text alternatives, code suggestions, and interface interaction. Reviews report growing use but uneven evaluation and limited participation by people with disabilities [8][10]. Lopez-Gil and Pereira reported 87.18 percent detection on selected WCAG checks that usually need manual judgment [9]. Yepes-Parra and Escobar-Velasquez also explored agents that interact through keyboard and screen-reader modes [11]. These studies support the direction of this work, but they do not settle how to produce safe, reviewable evidence for a complete journey.",
    )

    add_heading(doc, "2.3 Literature Comparison", 2)
    add_table(
        doc,
        ["Source", "Useful finding", "Limit for this project"],
        [
            ["WCAG 2.2 and WCAG-EM", "Define accessibility requirements and a review process", "Do not execute a complete user task"],
            ["Vigo et al. 2013", "Shows the risk of relying on automated tools alone", "Compares scanners rather than full journeys"],
            ["Ara et al. 2024", "Reviews 92 accessibility evaluation studies", "Finds broad methods and continuing gaps"],
            ["Pool 2023", "Shows that nine tools found different issues", "No tool provided complete coverage"],
            ["Lopez-Gil and Pereira 2025", "Uses an AI model for selected manual checks", "Covers three criteria, not task completion"],
            ["Aljedaani and Mollik 2026", "Reviews 38 studies of AI models and accessibility", "Finds varied methods and limited user involvement"],
            ["Yepes-Parra and Escobar-Velasquez 2025", "Explores agents using accessible interaction modes", "Early work with limited common benchmarks"],
        ],
        [1.55, 2.45, 2.7],
        font_size=8.7,
    )

    add_heading(doc, "3 Research Gap", 1)
    add_paragraph(
        doc,
        "AI has already been used in accessibility research, so this project does not claim otherwise. The narrower gap is the limited evidence on systems where an AI planner attempts a complete task through a declared accessibility method while fixed components control safety, execution, result classification, and evidence.",
    )
    add_bullets(
        doc,
        [
            "Page checks report isolated issues, while users succeed or fail across a sequence of states.",
            "Many tools inspect code or text without discovering and following a task path.",
            "AI can interpret intent, but it can also repeat, invent, or request unsafe actions.",
            "A useful barrier report needs the page state, focused element, attempted action, result, and screenshot.",
            "Controlled accessible and broken versions with known expected results remain uncommon.",
        ],
    )
    add_paragraph(
        doc,
        "Research gap statement: AccessPath AI studies a controlled framework in which an AI planner can select only from approved keyboard actions, while browser observation, safety checks, completion rules, and stored evidence remain under fixed program control.",
        bold_lead="Research gap statement:",
    )

    add_heading(doc, "4 Aim Objectives and Research Questions", 1)
    add_heading(doc, "4.1 Aim", 2)
    add_paragraph(
        doc,
        "To design and evaluate a safe, repeatable system that determines whether a multi-step web task can be completed by a keyboard-only user, using AI to suggest actions and fixed rules to control execution and results.",
    )

    doc.add_page_break()
    add_heading(doc, "4.2 Objectives", 2)
    add_numbered(
        doc,
        [
            "Build accessible and deliberately broken versions of the same task with known expected results.",
            "Use clear data structures for page observations, keyboard actions, safety decisions, and evidence.",
            "Support both a fixed test planner and selected live AI planners.",
            "Check every proposed action before it reaches the browser.",
            "Decide results from visible completion or barrier conditions, not from AI opinion.",
            "Store enough evidence to reproduce and review each result.",
            "Measure completion, correct blocking, false results, fallback use, and agreement across repeated runs.",
        ],
    )

    add_heading(doc, "4.3 Research Questions", 2)
    add_table(
        doc,
        ["ID", "Research question"],
        [
            ["RQ1", "Can a limited AI planner complete a multi-step task using only permitted keyboard actions?"],
            ["RQ2", "Can a full journey reveal keyboard barriers that a single-page scan does not explain well?"],
            ["RQ3", "Do fixed safety checks stop invalid or unsafe AI actions before browser execution?"],
            ["RQ4", "How reliably can the system separate completion, blocking, inconclusive results, and technical failure?"],
            ["RQ5", "Does step-by-step evidence make reported barriers easier to reproduce and review?"],
        ],
        [0.7, 6.0],
    )

    add_heading(doc, "5 Proposed Methodology", 1)
    add_heading(doc, "5.1 Research Design", 2)
    add_paragraph(
        doc,
        "The study follows a build-and-evaluate method. The team first creates a controlled system and test environment. It then runs the same declared task on accessible and broken versions, compares the result with the known expected result, and checks the stored evidence. AI selects a possible next action but does not decide safety or success.",
    )

    add_heading(doc, "5.2 Test Environment", 2)
    add_paragraph(
        doc,
        "Phase 1 uses versioned appointment-booking test sites. The accessible version supports keyboard completion. The broken version contains one declared barrier: the time-slot control cannot be reached or used through the permitted keyboard path. All input values are synthetic, so the study does not require patient or personal data. Using controlled sites makes the expected result known and repeated runs comparable.",
    )

    add_heading(doc, "5.3 System Architecture", 2)
    add_paragraph(
        doc,
        "The dashboard creates a run and sends it to a background worker. During the evaluation, the worker observes the page, obtains one proposed action, checks it, executes only an approved keyboard action, and applies fixed rules to the observed result. Run data and screenshots are saved for the reviewer.",
    )
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    add_figure(doc, DIAGRAM)

    add_heading(doc, "5.4 Evaluation Cycle", 2)
    add_table(
        doc,
        ["Step", "What happens"],
        [
            ["1 Observe", "Capture the approved URL, visible text, focused element, progress, and remaining limits."],
            ["2 Propose", "The selected planner suggests one permitted keyboard action."],
            ["3 Check", "Fixed rules verify the action, target, focus state, typing permission, and remaining limits."],
            ["4 Execute", "Playwright performs only the approved keyboard action."],
            ["5 Record", "Store the before and after states, result, screenshots, and planner details."],
            ["6 Decide", "Fixed rules confirm completion, a keyboard barrier, or a clearly labelled issue."],
        ],
        [1.1, 5.6],
    )

    add_heading(doc, "5.5 Tools and Technologies", 2)
    add_table(
        doc,
        ["Technology", "Purpose"],
        [
            ["Python and Django", "Dashboard, run workflow, validation, and result pages"],
            ["PostgreSQL", "Journeys, runs, steps, provider details, and evidence records"],
            ["Redis and Celery", "Background run queue and worker"],
            ["Playwright", "Keyboard input, browser observation, navigation, and screenshots"],
            ["Pydantic", "Checks that observations, actions, and evidence match the required structure"],
            ["Gemini and OpenRouter", "Optional live proposals; each provider is recorded separately"],
            ["Fixed test planner", "Repeatable controlled runs and visible contingency"],
            ["Docker Compose", "Reproducible local web, worker, database, queue, and test-site services"],
            ["pytest Ruff and Pyright", "Tests, code checks, and type checks"],
        ],
        [1.8, 4.9],
    )

    add_heading(doc, "5.6 Safety Controls", 2)
    add_bullets(
        doc,
        [
            "Only approved test-site addresses are allowed.",
            "The planner can request only a small set of keyboard actions; scripts and mouse clicks are not accepted.",
            "Fixed policy checks every proposal before execution.",
            "Text entry uses synthetic values only.",
            "Calls, steps, time, retries, and rejected actions all have limits.",
            "The report shows the provider and any fallback; fixed-plan evidence is never presented as live-AI evidence.",
            "Requests to the same provider are spaced out in the worker process.",
        ],
    )

    add_heading(doc, "6 Expected Outcomes and Evaluation", 1)
    add_heading(doc, "6.1 Outcome Classes", 2)
    add_bullets(
        doc,
        [
            "Completed: the declared confirmation state is visible.",
            "Blocked: evidence shows that a keyboard barrier prevents progress.",
            "Inconclusive: the run ends without enough evidence for completion or a confirmed barrier.",
            "Safety stopped: an action or navigation request breaks a safety rule.",
            "Provider or environment error: an external model or local service prevents evaluation.",
        ],
    )

    add_heading(doc, "6.2 Evaluation Measures", 2)
    add_table(
        doc,
        ["Measure", "Meaning"],
        [
            ["Task completion rate", "Share of accessible runs that reach the declared completion state"],
            ["Correct block rate", "Share of broken runs correctly reported as blocked"],
            ["False completion rate", "Broken runs incorrectly reported as completed; target is zero"],
            ["False block rate", "Accessible runs incorrectly reported as blocked"],
            ["Fallback rate", "Runs that switch from a live planner to the fixed contingency"],
            ["Repeated-run agreement", "Consistency of the result across repeated runs"],
            ["Evidence completeness", "Presence of action, focus, result, screenshot, reason, and provider details"],
            ["Efficiency", "Planner calls, executed steps, elapsed time, and provider failures"],
        ],
        [2.05, 4.65],
    )
    add_heading(doc, "7 Conclusion", 1)
    add_paragraph(
        doc,
        "AccessPath AI addresses the gap between finding accessibility issues on one page and showing whether a keyboard-only user can finish a complete task. AI can suggest the next action, but fixed rules control what reaches the browser, what counts as completion, and what evidence is stored. The current prototype demonstrates this approach on accessible and deliberately broken appointment-booking sites. Its contribution is a practical and reviewable way to study keyboard journey completion and barriers, not a claim that AI can certify accessibility.",
    )

    doc.add_page_break()
    add_heading(doc, "References", 1)
    references = [
        "World Wide Web Consortium. (2024). Web Content Accessibility Guidelines (WCAG) 2.2. https://www.w3.org/TR/WCAG22/",
        "World Wide Web Consortium. (2026). WCAG Evaluation Methodology (WCAG-EM) 2.0. https://www.w3.org/TR/WCAG-EM/",
        "Deque Systems. (2026). axe-core: Accessibility engine for automated Web UI testing. https://github.com/dequelabs/axe-core",
        "Microsoft. (2026). Accessibility testing with Playwright. https://playwright.dev/docs/accessibility-testing",
        "Vigo, M., Brown, J., and Conway, V. (2013). Benchmarking web accessibility evaluation tools: Measuring the harm of sole reliance on automated tests. W4A 2013. https://doi.org/10.1145/2461121.2461124",
        "Ara, J., Sik-Lanyi, C., and Kelemen, A. (2024). Accessibility engineering in web evaluation process: A systematic literature review. Universal Access in the Information Society, 23, 653-686. https://doi.org/10.1007/s10209-023-00967-2",
        "Pool, J. R. (2023). Accessibility Metatesting: Comparing Nine Testing Tools. W4A 2023. https://doi.org/10.1145/3587281.3587282",
        "Vera-Amaro, G., and Rojano-Caceres, J. R. (2025). Towards accessible website design through artificial intelligence: A systematic literature review. Information and Software Technology, 186, 107821. https://doi.org/10.1016/j.infsof.2025.107821",
        "Lopez-Gil, J. M., and Pereira, J. (2025). Turning manual web accessibility success criteria into automatic: An LLM-based approach. Universal Access in the Information Society, 24, 837-852. https://doi.org/10.1007/s10209-024-01108-z",
        "Aljedaani, W., and Mollik, R. H. (2026). Large Language Models for Web Accessibility: A Systematic Literature Review. W4A 2026. https://doi.org/10.1145/3800424.3800452",
        "Yepes-Parra, J. D., and Escobar-Velasquez, C. (2025). Autonomous Agents for Accessibility: Simulating Visual Impairments in Web Interfaces. ASE 2025, 3953-3957. https://doi.org/10.1109/ASE63991.2025.00349",
        "Microsoft. (2026). Keyboard API. Playwright documentation. https://playwright.dev/docs/api/class-keyboard",
    ]
    add_references(doc, references)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "AccessPath AI Project Synopsis"
    doc.core_properties.subject = "Keyboard journey accessibility evaluation"
    doc.core_properties.author = "Aniket Mishra, Krish Bajaj, Daksh Choudhary"
    doc.core_properties.keywords = "accessibility, keyboard, WCAG, AI planning, Playwright"
    doc.core_properties.comments = "Refined synopsis with project architecture diagram"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
